import time
import subprocess
import urllib.request
import urllib.parse
import csv
import codecs
import io
import os
import re
import unicodedata
from datetime import datetime

from PyQt6.QtCore import QThread, pyqtSignal

from utils import parse_data_jogo_bg
from config import URL_CSV_ANIVERSARIOS, URL_CSV_JOGO
from gerenciador_modulos import (
    alterar_estado_modulo,
    baixar_modulo_github,
    excluir_modulo,
    listar_modulos_github,
)


class WorkerModulosGitHub(QThread):
    resultado = pyqtSignal(object)
    erro = pyqtSignal(str)

    def __init__(self, nome=None, url=None, acao=None, parent=None):
        super().__init__(parent)
        self.nome = nome
        self.url = url
        self.acao = acao

    def run(self):
        try:
            if self.nome is None:
                resultado = listar_modulos_github()
            elif self.acao == "baixar":
                resultado = baixar_modulo_github(self.nome, self.url)
            elif self.acao == "ativar":
                resultado = alterar_estado_modulo(self.nome, True)
            elif self.acao == "desativar":
                resultado = alterar_estado_modulo(self.nome, False)
            elif self.acao == "excluir":
                resultado = excluir_modulo(self.nome)
            else:
                raise ValueError("Ação de módulo inválida.")
            if self.nome is not None:
                resultado = {"nome": resultado, "acao": self.acao}
            self.resultado.emit(resultado)
        except Exception as erro:
            self.erro.emit(str(erro))


class WorkerDownload(QThread):
    resultado = pyqtSignal(list, list, object, object, bytes, bytes)

    def __init__(self, planilha_arquivo="", planilha_url="", parent=None):
        super().__init__(parent)
        self.planilha_arquivo = planilha_arquivo or ""
        self.planilha_url = planilha_url or ""
        self.erro_planilha = ""

    @staticmethod
    def _normalizar_texto(valor):
        if valor is None:
            return ""
        texto = str(valor).strip()
        if not texto:
            return ""
        texto = unicodedata.normalize("NFKC", texto)
        return texto

    @staticmethod
    def _normalizar_cabecalho(texto):
        texto = WorkerDownload._normalizar_texto(texto).casefold()
        texto = re.sub(r"[^a-z0-9]+", "", texto)
        return texto

    @staticmethod
    def _detectar_delimitador(texto):
        amostra = texto[:4096]
        try:
            return csv.Sniffer().sniff(amostra, delimiters=",;\t|").delimiter
        except Exception:
            if amostra.count(";") > amostra.count(","):
                return ";"
            return ","

    @staticmethod
    def _url_google_sheets_para_csv(url):
        partes = urllib.parse.urlparse(url)
        if "docs.google.com" not in partes.netloc:
            return url
        if "/spreadsheets/d/" not in partes.path:
            return url
        match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", partes.path)
        if not match:
            return url
        spreadsheet_id = match.group(1)
        query = urllib.parse.parse_qs(partes.query)
        gid = query.get("gid", ["0"])[0]
        return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=csv&gid={gid}"

    @staticmethod
    def _eh_html(bruto_texto):
        inicio = bruto_texto.lstrip().lower()
        return inicio.startswith("<!doctype html") or inicio.startswith("<html") or "<body" in inicio[:500]

    def _normalizar_linhas(self, linhas):
        linhas = list(linhas)
        if not linhas:
            return []

        indice_nome = 0
        indice_data = 1
        indice_cor = 2
        encontrou_cabecalho = False
        cabecalho = [self._normalizar_cabecalho(c) for c in linhas[0]]
        mapa = {valor: idx for idx, valor in enumerate(cabecalho) if valor}
        if mapa:
            for nomes, destino in (
                ({"nome", "nomecompleto", "pessoa", "aniversariante"}, "nome"),
                ({"data", "dataaniversario", "aniversario", "nascimento", "datanascimento"}, "data"),
                ({"cor", "corhex", "hex", "color"}, "cor"),
            ):
                for candidato in nomes:
                    if candidato in mapa:
                        if destino == "nome":
                            indice_nome = mapa[candidato]
                            encontrou_cabecalho = True
                        elif destino == "data":
                            indice_data = mapa[candidato]
                            encontrou_cabecalho = True
                        elif destino == "cor":
                            indice_cor = mapa[candidato]
                            encontrou_cabecalho = True
                        break

        inicio_linhas = 1 if encontrou_cabecalho else 0
        normalizados = []
        for linha in linhas[inicio_linhas:]:
            if not linha:
                continue
            if len(linha) <= max(indice_nome, indice_data):
                continue
            nome = self._normalizar_texto(linha[indice_nome])
            data = self._normalizar_texto(linha[indice_data])
            if not nome or not data:
                continue
            cor = "#ffffff"
            if len(linha) > indice_cor:
                cor_bruta = self._normalizar_texto(linha[indice_cor])
                if cor_bruta:
                    cor = cor_bruta
            normalizados.append((nome, data, cor))
        return normalizados

    def _ler_arquivo_planilha(self, caminho):
        if not os.path.exists(caminho):
            raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

        extensao = os.path.splitext(caminho)[1].lower()
        if extensao in (".csv", ".txt", ".tsv"):
            with open(caminho, "r", encoding="utf-8-sig", errors="replace") as arquivo:
                texto = arquivo.read()
            delimitador = self._detectar_delimitador(texto)
            return self._normalizar_linhas(csv.reader(io.StringIO(texto), delimiter=delimitador))

        if extensao in (".xlsx", ".xlsm"):
            try:
                from openpyxl import load_workbook
            except Exception as erro:
                raise RuntimeError("openpyxl não está disponível para ler planilhas .xlsx.") from erro

            workbook = load_workbook(caminho, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                return self._normalizar_linhas(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()

        if extensao == ".xls":
            try:
                import xlrd
            except Exception as erro:
                raise RuntimeError("xlrd não está disponível para ler planilhas .xls.") from erro

            workbook = xlrd.open_workbook(caminho)
            sheet = workbook.sheet_by_index(0)
            linhas = (sheet.row_values(i) for i in range(sheet.nrows))
            return self._normalizar_linhas(linhas)

        raise ValueError(f"Formato de planilha não suportado: {extensao or caminho}")

    def _ler_url_planilha(self, url):
        url = self._url_google_sheets_para_csv(self._normalizar_texto(url))
        if not url:
            raise ValueError("URL da planilha vazia.")

        with urllib.request.urlopen(url, timeout=15) as resposta:
            conteudo = resposta.read()

        texto = None
        for encoding in ("utf-8-sig", "utf-8", "cp1252"):
            try:
                texto = conteudo.decode(encoding)
                break
            except Exception:
                continue
        if texto is None:
            raise ValueError("Não foi possível decodificar o conteúdo da planilha.")

        if self._eh_html(texto):
            raise ValueError("O link retornou uma página HTML em vez da planilha CSV.")

        delimitador = self._detectar_delimitador(texto)
        return self._normalizar_linhas(csv.reader(io.StringIO(texto), delimiter=delimitador))

    def run(self):
        dados_jogo = [{"imagem": "", "link": ""}, {"imagem": "", "link": ""}]
        data_inicio = None
        data_fim = None
        img_bytes1 = b""
        img_bytes2 = b""
        combinados = []
        vistos = set()
        self.erro_planilha = ""

        def registrar_linhas(linhas):
            for nome, data, cor in linhas:
                chave = (nome.casefold(), data, cor.casefold())
                if chave in vistos:
                    continue
                vistos.add(chave)
                combinados.append((nome, data, cor))

        fontes = []
        if self.planilha_arquivo:
            fontes.append((self.planilha_arquivo, True, self._ler_arquivo_planilha))
        if self.planilha_url:
            fontes.append((self.planilha_url, True, self._ler_url_planilha))
        fontes.append((URL_CSV_ANIVERSARIOS, False, self._ler_url_planilha))

        for origem, obrigatoria, carregador in fontes:
            try:
                registrar_linhas(carregador(origem))
            except Exception as erro:
                if obrigatoria and not self.erro_planilha:
                    self.erro_planilha = str(erro)

        if not combinados:
            self.erro_planilha = self.erro_planilha or "Nenhuma fonte de aniversários retornou dados."

        try:
            req_jogo = urllib.request.urlopen(URL_CSV_JOGO, timeout=10)
            csv_jogo = list(csv.reader(codecs.iterdecode(req_jogo, "utf-8")))
            if len(csv_jogo) > 1:
                linha2 = csv_jogo[1]
                if len(linha2) >= 1:
                    dados_jogo[0]["imagem"] = linha2[0].strip()
                if len(linha2) >= 2:
                    dados_jogo[0]["link"] = linha2[1].strip()
                if len(linha2) >= 4:
                    data_inicio = parse_data_jogo_bg(linha2[2].strip())
                    data_fim = parse_data_jogo_bg(linha2[3].strip())
            if len(csv_jogo) > 2:
                linha3 = csv_jogo[2]
                if len(linha3) >= 1:
                    dados_jogo[1]["imagem"] = linha3[0].strip()
                if len(linha3) >= 2:
                    dados_jogo[1]["link"] = linha3[1].strip()
        except Exception:
            pass

        try:
            if dados_jogo[0]["imagem"]:
                req = urllib.request.Request(dados_jogo[0]["imagem"], headers={"User-Agent": "Mozilla/5.0"})
                img_bytes1 = urllib.request.urlopen(req, timeout=10).read()
        except Exception:
            pass

        try:
            if dados_jogo[1]["imagem"]:
                req = urllib.request.Request(dados_jogo[1]["imagem"], headers={"User-Agent": "Mozilla/5.0"})
                img_bytes2 = urllib.request.urlopen(req, timeout=10).read()
        except Exception:
            pass

        self.resultado.emit(combinados, dados_jogo, data_inicio, data_fim, img_bytes1, img_bytes2)


class WorkerMonitorProcessos(QThread):
    processos_atualizados = pyqtSignal(set)

    def __init__(self, parent=None):
        super().__init__(parent)
        self._executando = True

    def parar(self):
        self._executando = False

    def run(self):
        CREATE_NO_WINDOW = 0x08000000
        while self._executando:
            inicio = time.monotonic()
            try:
                out = subprocess.check_output(
                    ["tasklist", "/FO", "CSV", "/NH"],
                    creationflags=CREATE_NO_WINDOW,
                    text=True,
                )
                nomes = set(
                    line.split(",")[0].strip('"').lower()
                    for line in out.splitlines() if line
                )
                self.processos_atualizados.emit(nomes)
            except Exception:
                pass

            time.sleep(max(0.05, 0.50 - (time.monotonic() - inicio)))
